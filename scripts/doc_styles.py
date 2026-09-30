# -*- coding: utf-8 -*-
"""
Script para actualizar y completar el documento 'Docs/ALBUM UML Bloque3_YAPU.docx'
incorporando el Documento de Gestión completo según los requerimientos académicos:
- Portada institucional UPDS
- Resumen ejecutivo integrado
- Introducción contextualizada
- Marco Teórico ampliado (Pirámide de pruebas, Viabilidad, Costo-beneficio socioambiental, Riesgos de IA, Gobierno)
- Metodología de gestión y aseguramiento de la calidad
- Resultados: Modelos UML base + 6 secciones del Documento de Gestión:
  1. Plan de pruebas
  2. Suite de pruebas ejecutadas (630 tests, 100% verde, V8 93.56%, capturas reales)
  3. Estudio de viabilidad multidimensional
  4. Análisis costo/beneficio con dimensión socioambiental
  5. Matriz de riesgos del proyecto y riesgos de IA (ISO 31000 & OWASP)
  6. Tablero de métricas de gobierno de software
- Discusión académica ampliada
- Conclusiones integrales
- Referencias en APA 7 con sangría francesa
- Anexos A, B, C, D conservados y ampliados + Anexo E (registros de pruebas) + Anexo F (álbum fotográfico en vivo)
"""

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn
import os
import shutil

DOC_PATH = 'Docs/ALBUM UML Bloque3_YAPU.docx'
BACKUP_PATH = 'Docs/ALBUM UML Bloque3_YAPU_original_backup.docx'

def aplicar_estilos_parrafo(p, font_name="Times New Roman", size_pt=12, line_spacing=1.5, space_after=6, space_before=0, align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p.paragraph_format.line_spacing = line_spacing
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.alignment = align
    for run in p.runs:
        run.font.name = font_name
        run.font.size = Pt(size_pt)

def set_cell_shading(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=180, right=180):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="8" w:space="0" w:color="334155"/>'
        f'<w:bottom w:val="single" w:sz="8" w:space="0" w:color="334155"/>'
        f'<w:insideH w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>'
        f'<w:insideV w:val="none"/>'
        f'<w:left w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def crear_tabla_estilizada(doc, headers, data, col_widths=None):
    tbl = doc.add_table(rows=len(data) + 1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl)
    
    # Encabezado
    hdr_cells = tbl.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_shading(hdr_cells[i], "1E293B")
        set_cell_margins(hdr_cells[i], top=140, bottom=140, left=160, right=160)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.space_before = Pt(2)
        for r in p.runs:
            r.font.name = "Times New Roman"
            r.font.size = Pt(10)
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)
            
    # Filas
    for row_idx, row_data in enumerate(data):
        row_cells = tbl.rows[row_idx + 1].cells
        bg_col = "F8FAFC" if row_idx % 2 == 1 else "FFFFFF"
        for col_idx, cell_value in enumerate(row_data):
            row_cells[col_idx].text = str(cell_value)
            set_cell_shading(row_cells[col_idx], bg_col)
            set_cell_margins(row_cells[col_idx], top=100, bottom=100, left=140, right=140)
            p = row_cells[col_idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.space_before = Pt(2)
            for r in p.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(9.5)
                
    if col_widths:
        for row in tbl.rows:
            for idx, width in enumerate(col_widths):
                row.cells[idx].width = Inches(width)
                
    return tbl

print("Módulo de estilos cargado exitosamente.")
