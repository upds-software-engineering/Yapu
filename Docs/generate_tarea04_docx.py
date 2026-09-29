# Script para compilar Informe de Tarea 4 en DOCX (APA 7) y PDF
import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_shading(cell, color_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def style_apa_table(table):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for r_idx, row in enumerate(table.rows):
        for cell in row.cells:
            tcPr = cell._tc.get_or_add_tcPr()
            tcBorders = OxmlElement('w:tcBorders')
            
            top = OxmlElement('w:top')
            if r_idx == 0:
                top.set(qn('w:val'), 'single')
                top.set(qn('w:sz'), '12')
                top.set(qn('w:space'), '0')
                top.set(qn('w:color'), '000000')
            else:
                top.set(qn('w:val'), 'none')
            tcBorders.append(top)
            
            bottom = OxmlElement('w:bottom')
            if r_idx == 0:
                bottom.set(qn('w:val'), 'single')
                bottom.set(qn('w:sz'), '6')
                bottom.set(qn('w:space'), '0')
                bottom.set(qn('w:color'), '000000')
            elif r_idx == len(table.rows) - 1:
                bottom.set(qn('w:val'), 'single')
                bottom.set(qn('w:sz'), '12')
                bottom.set(qn('w:space'), '0')
                bottom.set(qn('w:color'), '000000')
            else:
                bottom.set(qn('w:val'), 'none')
            tcBorders.append(bottom)
            
            for side in ['left', 'right']:
                b = OxmlElement(f'w:{side}')
                b.set(qn('w:val'), 'none')
                tcBorders.append(b)
                
            tcPr.append(tcBorders)
            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)

def build_docx_from_md(md_path, docx_path):
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    doc = docx.Document()

    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)

    style_normal = doc.styles['Normal']
    font_normal = style_normal.font
    font_normal.name = 'Times New Roman'
    font_normal.size = Pt(12)
    font_normal.color.rgb = RGBColor(0, 0, 0)

    def add_p(text="", bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=6, space_before=0, line_spacing=1.5):
        p = doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.line_spacing = line_spacing
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.space_before = Pt(space_before)
        if text:
            run = p.add_run(text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(12)
            run.font.bold = bold
            run.font.italic = italic
        return p

    def add_code_block(code_lines, border_color='0284C7', bg_color='F8FAFC'):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        cell = tbl.cell(0, 0)
        set_cell_shading(cell, bg_color)
        set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
        
        tcPr = cell._tc.get_or_add_tcPr()
        tcBorders = OxmlElement('w:tcBorders')
        left = OxmlElement('w:left')
        left.set(qn('w:val'), 'single')
        left.set(qn('w:sz'), '16')
        left.set(qn('w:color'), border_color)
        tcBorders.append(left)
        for s in ['top', 'bottom', 'right']:
            b = OxmlElement(f'w:{s}')
            b.set(qn('w:val'), 'none')
            tcBorders.append(b)
        tcPr.append(tcBorders)

        cp = cell.paragraphs[0]
        cp.paragraph_format.line_spacing = 1.15
        cp.paragraph_format.space_after = Pt(2)
        cp.paragraph_format.space_before = Pt(2)
        for c_idx, line in enumerate(code_lines):
            r = cp.add_run(line + ('\n' if c_idx < len(code_lines) - 1 else ''))
            r.font.name = 'Consolas'
            r.font.size = Pt(10)
            r.font.color.rgb = RGBColor(30, 41, 59)

    in_code = False
    code_buffer = []
    in_table = False
    table_buffer = []

    for raw_line in lines:
        line = raw_line.rstrip('\r\n')

        # Code block handling
        if line.startswith('```'):
            if in_code:
                in_code = False
                add_code_block(code_buffer)
                code_buffer = []
            else:
                in_code = True
                code_buffer = []
            continue

        if in_code:
            code_buffer.append(line)
            continue

        # Table handling
        if '|' in line and not line.startswith('```'):
            if re.match(r'^\s*\|?\s*[-:]+[-| :]*$', line):
                # separator row
                continue
            cols = [c.strip() for c in line.split('|')]
            if cols and cols[0] == '': cols.pop(0)
            if cols and cols[-1] == '': cols.pop()
            if cols:
                in_table = True
                table_buffer.append(cols)
                continue
        else:
            if in_table and table_buffer:
                # Render table
                num_cols = max(len(row) for row in table_buffer)
                tbl = doc.add_table(rows=len(table_buffer), cols=num_cols)
                for r_idx, row in enumerate(table_buffer):
                    for c_idx, cell_text in enumerate(row):
                        if c_idx < num_cols:
                            cell = tbl.cell(r_idx, c_idx)
                            clean_text = re.sub(r'\*\*(.*?)\*\*', r'\1', cell_text)
                            clean_text = re.sub(r'\*(.*?)\*', r'\1', clean_text)
                            clean_text = re.sub(r'`(.*?)`', r'\1', clean_text)
                            cp = cell.paragraphs[0]
                            cp.paragraph_format.line_spacing = 1.15
                            cp.paragraph_format.space_after = Pt(2)
                            cp.paragraph_format.space_before = Pt(2)
                            r = cp.add_run(clean_text)
                            r.font.name = 'Times New Roman'
                            r.font.size = Pt(10)
                            if r_idx == 0:
                                r.font.bold = True
                                set_cell_shading(cell, "F1F5F9")
                style_apa_table(tbl)
                add_p("", space_after=6)
                in_table = False
                table_buffer = []

        # Headings
        if line.startswith('# '):
            text = line[2:].strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(8)
            p.paragraph_format.line_spacing = 1.5
            r = p.add_run(text)
            r.font.name = 'Times New Roman'
            r.font.size = Pt(15)
            r.font.bold = True
            r.font.color.rgb = RGBColor(15, 23, 42)
            continue

        if line.startswith('## '):
            text = line[3:].strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.5
            r = p.add_run(text)
            r.font.name = 'Times New Roman'
            r.font.size = Pt(13)
            r.font.bold = True
            r.font.color.rgb = RGBColor(30, 41, 59)
            continue

        if line.startswith('### '):
            text = line[4:].strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.5
            r = p.add_run(text)
            r.font.name = 'Times New Roman'
            r.font.size = Pt(12)
            r.font.bold = True
            r.font.italic = True
            r.font.color.rgb = RGBColor(51, 65, 85)
            continue

        if line.startswith('---'):
            # Horizontal rule
            continue

        if line.strip() == '':
            continue

        # Regular bullet or text
        is_bullet = False
        prefix = ""
        content = line
        if line.startswith('* ') or line.startswith('- '):
            is_bullet = True
            content = line[2:].strip()
        elif re.match(r'^\d+\.\s+', line):
            is_bullet = False
            m = re.match(r'^(\d+\.\s+)(.*)$', line)
            prefix = m.group(1)
            content = m.group(2).strip()

        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.space_before = Pt(0)
        if is_bullet:
            p.paragraph_format.left_indent = Inches(0.25)
            r_bullet = p.add_run("• ")
            r_bullet.font.name = 'Times New Roman'
            r_bullet.font.size = Pt(12)
            r_bullet.font.bold = True
        elif prefix:
            r_pref = p.add_run(prefix)
            r_pref.font.name = 'Times New Roman'
            r_pref.font.size = Pt(12)
            r_pref.font.bold = True

        # Parse inline markdown formatting (bold, italic, code)
        tokens = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)', content)
        for token in tokens:
            if not token:
                continue
            if token.startswith('**') and token.endswith('**'):
                r = p.add_run(token[2:-2])
                r.font.name = 'Times New Roman'
                r.font.size = Pt(12)
                r.font.bold = True
            elif token.startswith('*') and token.endswith('*'):
                r = p.add_run(token[1:-1])
                r.font.name = 'Times New Roman'
                r.font.size = Pt(12)
                r.font.italic = True
            elif token.startswith('`') and token.endswith('`'):
                r = p.add_run(token[1:-1])
                r.font.name = 'Consolas'
                r.font.size = Pt(10.5)
                r.font.color.rgb = RGBColor(194, 65, 12)
            else:
                r = p.add_run(token)
                r.font.name = 'Times New Roman'
                r.font.size = Pt(12)

    doc.save(docx_path)
    print(f"DOCX creado exitosamente en: {docx_path}")

def convert_to_pdf(docx_path, pdf_path):
    import win32com.client
    import pythoncom
    pythoncom.CoInitialize()
    word = win32com.client.DispatchEx("Word.Application")
    word.Visible = False
    word.DisplayAlerts = False
    try:
        abs_docx = os.path.abspath(docx_path)
        abs_pdf = os.path.abspath(pdf_path)
        doc = word.Documents.Open(abs_docx)
        doc.SaveAs(abs_pdf, FileFormat=17) # 17 = wdFormatPDF
        doc.Close()
        print(f"PDF compilado exitosamente en: {pdf_path}")
    finally:
        word.Quit()
        pythoncom.CoUninitialize()

if __name__ == '__main__':
    base_dir = r"J:\www\Yapu\Docs"
    md_file = os.path.join(base_dir, "Informe_Tarea04_Gestion_Gobernanza_Riesgos.md")
    docx_file = os.path.join(base_dir, "Informe_Tarea04_Gestion_Gobernanza_Riesgos.docx")
    pdf_file = os.path.join(base_dir, "Informe_Tarea04_Gestion_Gobernanza_Riesgos.pdf")

    build_docx_from_md(md_file, docx_file)
    try:
        convert_to_pdf(docx_file, pdf_file)
    except Exception as e:
        print(f"Aviso al compilar PDF: {e}")
